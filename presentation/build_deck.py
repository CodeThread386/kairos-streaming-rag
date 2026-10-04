"""Fills the official CollegeName_TeamName_Submission.pptx template (12 slides) with Kairos content.
Run:  python presentation/build_deck.py <template.pptx> <screens_dir> <out.pptx>"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

TPL, SHOTS, OUT = sys.argv[1], Path(sys.argv[2]), sys.argv[3]
ROOT = Path(__file__).resolve().parents[1]

PURPLE = RGBColor(0x70, 0x4E, 0xA6)
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
INK = RGBColor(0x14, 0x14, 0x2B)
MUTED = RGBColor(0x63, 0x63, 0x7E)
OK = RGBColor(0x13, 0x8A, 0x62)
BLUE = RGBColor(0x33, 0x55, 0xFF)
TINT = RGBColor(0xF3, 0xF0, 0xFA)
RULE = RGBColor(0xDD, 0xD8, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation(TPL)
S = list(prs.slides)
W, H = prs.slide_width, prs.slide_height
LM = Inches(0.92)                 # template content left margin
CW = W - 2 * LM


def drop_body(slide):
    for sh in list(slide.placeholders):
        if sh.placeholder_format.type is not None and sh.placeholder_format.idx == 1:
            sh._element.getparent().remove(sh._element)


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.0, name=None):
    """runs: str | list[paragraph], paragraph = str | list[(text, {bold,color,size})]"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = para if isinstance(para, list) else [(para, {})]
        for t, opt in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = opt.get("color", color)
        if i < len(paras) - 1:
            p.space_after = Pt(opt.get("after", 6) if parts else 6)
    return tb


def card(slide, x, y, w, h, fill=TINT, line=None, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if radius:
        shp.adjustments[0] = 0.08
    return shp


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.22))))
        pPr.set("indent", str(-Emu(Inches(0.22))))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        p.space_after = Pt(gap)
        parts = it if isinstance(it, list) else [(it, {})]
        for t, opt in parts:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", False)
            r.font.color.rgb = opt.get("color", color)
    return tb


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), x, y, w, h)


def table(slide, x, y, w, rows, col_w, size=13, header_fill=PURPLE, row_h=Inches(0.42), bold_last=False):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = cw
    for i, row in enumerate(rows):
        tbl.rows[i].height = row_h
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (WHITE if i % 2 else TINT)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = i == 0 or (bold_last and j == len(row) - 1)
            r.font.color.rgb = WHITE if i == 0 else (OK if (bold_last and j == len(row) - 1) else INK)
    return shp


# ---------------------------------------------------------------- slide 1: title / team details
s1 = S[0]
info = next(sh for sh in s1.shapes if sh.has_text_frame and "Theme ID" in sh.text_frame.text)
lines = ["Theme ID - 04 · Streaming Live RAG",
         "Project - Kairos: streaming live RAG that retrieves while you speak",
         "Team Name - VITV_Kairos_4",
         "College Name - VIT Vellore (VITV)",
         "Member Name & Email 1- Sarthak Trivedi, sarthak.trivedi2024@vitstudent.ac.in",
         "Member Name 2- Gurutheja Kavuturi",
         "Member Name 3- Nigam Kundu",
         "Submission Github link - github.com/CodeThread386/kairos-streaming-rag (tag PRISM_GENAI_HACKATHON_Y2026)"]
paras = info.text_frame.paragraphs
template_p = copy.deepcopy(paras[0]._p)
txBody = info.text_frame._txBody
for p in list(paras):
    txBody.remove(p._p)
for ln in lines:
    np_ = copy.deepcopy(template_p)
    rs = np_.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
    for extra in rs[1:]:
        np_.remove(extra)
    rs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = ln
    txBody.append(np_)
info.height = Emu(int(info.height * 2.2))

# ---------------------------------------------------------------- slide 2: theme / problem
s = S[1]
drop_body(s)
text(s, LM, Inches(1.55), Inches(7.2), Inches(0.9),
     [[("Theme 04, Streaming Live RAG: ", {"bold": True, "color": PURPLE}),
       ("in a full-duplex voice conversation, retrieve the right context mid-conversation from one natural command.", {})]],
     size=18)
probs = [("It waits for silence", "Batch RAG does nothing while the user speaks, then pays the whole retrieval round trip after they stop."),
         ("One query for many asks", "\"Venues for 30 in Pune, the cancellation policy and the catering options\" is three questions. One embedding blurs them."),
         ("Late details restart it", "\"Actually, the trip was international.\" Restarting discards evidence, costs another full search and loses citations.")]
for i, (h_, b_) in enumerate(probs):
    y = Inches(2.65) + i * Inches(1.18)
    card(s, LM, y, Inches(7.2), Inches(1.02))
    text(s, LM + Inches(0.25), y + Inches(0.14), Inches(6.7), Inches(0.8), [[(h_, {"bold": True, "size": 17})], [(b_, {"color": MUTED, "size": 14})]])
text(s, Inches(8.5), Inches(1.6), Inches(4.0), Inches(0.5), [[("What we committed to", {"bold": True, "color": PURPLE})]], size=18)
bullets(s, Inches(8.5), Inches(2.2), Inches(4.0), Inches(4.2), [
    [("Listen incrementally: ", {"bold": True}), ("decide on every chunk whether to wait, retrieve or skip", {})],
    [("Decompose ", {"bold": True}), ("one utterance into the questions it actually implies", {})],
    [("Refine, never restart, ", {"bold": True}), ("when a constraint arrives late", {})],
    [("Ground every claim ", {"bold": True}), ("in a cited section, or say it is unverified", {})],
    [("Prove it: ", {"bold": True}), ("measure all six acceptance gates G1–G6 on held-out streams", {})]], size=15)

# ---------------------------------------------------------------- slide 3: existing solutions & gaps
s = S[2]
drop_body(s)
rows = [["Approach", "How it works", "Gap for live voice"],
        ["Batch RAG (default chains in common RAG frameworks)", "Waits for end of utterance, embeds it as one query, retrieves top-k", "Retrieval starts only after the user stops; compound requests become one blurred query"],
        ["LLM query rewriting / multi-query retrieval", "An LLM rewrites or splits the question before searching", "Adds an LLM round trip and cost per turn, still after the utterance ends"],
        ["Realtime speech LLM APIs", "Stream audio in and out with low latency", "Tool or retrieval calls fire after the turn; grounding and citations are left to the app"],
        ["Conversational RAG with history", "Condenses chat history into a standalone query each turn", "A late detail triggers a full new search; earlier evidence and citations are not carried forward"]]
table(s, LM, Inches(1.65), CW, rows, [Inches(3.1), Inches(4.0), CW - Inches(7.1)], size=14, row_h=Inches(0.78))
card(s, LM, Inches(5.75), CW, Inches(0.78), fill=TINT)
text(s, LM + Inches(0.3), Inches(5.85), CW - Inches(0.6), Inches(0.6),
     [[("The gap: ", {"bold": True, "color": PURPLE}),
       ("nothing decides when to retrieve while the user is still talking, splits the request cheaply, and patches a cited answer instead of restarting. Kairos does all three on a CPU, with no LLM in the loop.", {})]],
     size=15, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- slide 4: solution & architecture
s = S[3]
drop_body(s)
ph = Inches(5.45)
pw_ = int(ph * 1600 / 900)
pic = picture(s, ROOT / "docs" / "architecture.png", int((W - pw_) / 2), Inches(1.4), h=ph)
text(s, LM, Inches(1.4) + pic.height + Inches(0.1), CW, Inches(0.4),
     "Controller ≈ 1 ms per chunk · retrieval ≈ 0.6 ms per query · answer starts ≈ 2 ms after the user stops · 0 LLM calls by default",
     size=13, color=MUTED, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- slide 5: demo walkthrough
s = S[4]
drop_body(s)
pw = (CW - Inches(0.3)) / 2
for i, (img, cap) in enumerate([("ui_workshop.png", "Three questions in one breath: retrievals fire 6.4 s before the speaker stops; answer starts 14 ms after."),
                                ("ui_refine.png", "Late detail \"international, booked after travel\": only the new details are searched; v2 keeps v1's citations.")]):
    x = LM + i * (pw + Inches(0.3))
    pic = picture(s, SHOTS / img, x, Inches(1.45), w=pw)
    pic.line.color.rgb = RULE
    text(s, x, Inches(1.45) + pic.height + Inches(0.08), pw, Inches(0.5), cap, size=12, color=MUTED)
steps = [("Speech timeline", "Words as spoken; blue markers are retrievals fired before the amber \"you stopped\" line"),
         ("Controller log", "Every WAIT / RETRIEVE / SUPPRESS decision with its reason and each sub-query's hits"),
         ("Versioned, cited answer", "Citation chips open the source sentence; v2 lists kept, added and replaced claims"),
         ("Live telemetry", "Time to first token, retrieval lead, queries fired and used, reuse, cost")]
cw4 = (CW - Inches(0.9)) / 4
for i, (h_, b_) in enumerate(steps):
    x = LM + i * (cw4 + Inches(0.3))
    y = Inches(5.25)
    card(s, x, y, cw4, Inches(1.35))
    text(s, x + Inches(0.18), y + Inches(0.12), cw4 - Inches(0.36), Inches(1.15),
         [[(h_, {"bold": True, "size": 14, "color": PURPLE})], [(b_, {"size": 12})]])
text(s, LM, Inches(6.72), CW, Inches(0.3), [[("Demo video (3:30): ", {"bold": True}), ("demo/kairos_demo.mp4 in the repository; YouTube link in the README", {})]], size=12, color=MUTED)

# ---------------------------------------------------------------- slide 6: tools & tech stack
s = S[5]
drop_body(s)
groups = [("Streaming core", "Python 3.12 · asyncio · FastAPI + WebSocket · Uvicorn"),
          ("Retrieval", "BM25 (own implementation) · WordLlama 256-d static embeddings (CPU, ships in the wheel) · reciprocal-rank fusion · NumPy"),
          ("Language processing", "Snowball stemmer · corpus-derived lexicons (IDF, entities, scopes, aspects) · rule-based self-repair and clause segmentation"),
          ("Ingestion", "Markdown / HTML / TXT / JSON parsers · pypdf · python-docx → section chunks with Doc_ID §N"),
          ("Interface", "Single-page console (HTML + JS) · Web Speech API for live mic input · typed input streamed word by word"),
          ("Quality & ops", "pytest (16 tests) · jsonschema telemetry validation · replay benchmark harness · Docker + compose · GitHub Actions CI"),
          ("Optional LLM", "Any OpenAI-compatible endpoint (Ollama, vLLM, OpenAI, Gemini, Groq) with a sentence-level grounding verifier"),
          ("Demo production", "Playwright (scripted UI recording) · Kokoro-82M TTS via sherpa-onnx (offline narration) · ffmpeg")]
cw_ = (CW - Inches(0.3)) / 2
for i, (h_, b_) in enumerate(groups):
    col, row = i % 2, i // 2
    x = LM + col * (cw_ + Inches(0.3))
    y = Inches(1.5) + row * Inches(1.25)
    card(s, x, y, cw_, Inches(1.1))
    text(s, x + Inches(0.22), y + Inches(0.13), cw_ - Inches(0.44), Inches(0.9),
         [[(h_, {"bold": True, "size": 16, "color": PURPLE})], [(b_, {"size": 13})]])

# ---------------------------------------------------------------- slide 7: impact & use case
s = S[6]
drop_body(s)
uses = [("Galaxy voice assistant", "\"Turn on Wi-Fi calling and how do I stop apps draining battery\" retrieved from device manuals while the user is still speaking."),
        ("Enterprise help desks", "HR, travel and IT policy questions with several asks per sentence, answered with citations auditors can check."),
        ("In-car and hands-free", "Drivers add details mid-request (\"actually, the route via the highway\"). The answer is patched, not restarted."),
        ("Live agent assist", "Surfaces cited policy snippets to a call-centre agent before the customer finishes the sentence.")]
for i, (h_, b_) in enumerate(uses):
    y = Inches(1.5) + i * Inches(1.18)
    card(s, LM, y, Inches(7.4), Inches(1.03))
    text(s, LM + Inches(0.25), y + Inches(0.13), Inches(6.9), Inches(0.85), [[(h_, {"bold": True, "size": 16})], [(b_, {"size": 13, "color": MUTED})]])
stats = [("2.4 ms", "time to first token after the user stops, even when each search takes 150 ms (batch RAG: 152 ms)"),
         ("+27 pts", "answers that cite the right section, vs batch RAG (62.7% → 89.5%)"),
         ("$0 / turn", "no LLM in the loop; 16 MB embedding model, CPU only, so it fits on device")]
for i, (big, small) in enumerate(stats):
    y = Inches(1.5) + i * Inches(1.6)
    text(s, Inches(8.75), y, Inches(3.7), Inches(1.5), [[(big, {"bold": True, "size": 36, "color": VIOLET})], [(small, {"size": 13, "color": MUTED})]])

# ---------------------------------------------------------------- slide 8: innovation, results, limitations
s = S[7]
drop_body(s)
text(s, LM, Inches(1.45), Inches(5.4), Inches(0.4), [[("Innovation", {"bold": True, "color": PURPLE})]], size=17)
bullets(s, LM, Inches(1.9), Inches(5.4), Inches(2.9), [
    "Retrieval controller with dangling-phrase and corpus-IDF stability checks: early retrieval with 0% false triggers",
    "Speculative retrieval reused only when it covers every final word (93–94% reuse)",
    "Weak context propagation and entity-conditional scoping across sub-queries",
    "Versioned answers: preserved, added and superseded claims, plus zero-query reformatting",
    "Extractive grounding by construction, plus a verifier that keeps an optional LLM honest"], size=13, gap=5)
text(s, LM, Inches(4.95), Inches(5.4), Inches(0.4), [[("Limitations", {"bold": True, "color": PURPLE})]], size=17)
bullets(s, LM, Inches(5.4), Inches(5.4), Inches(1.4), [
    "Synonym gaps (\"cap\" vs \"limit\") cause 4–6% false \"unverified\" flags",
    "Multi-hop joins (city → tier → rate) are cited, not computed",
    "Harder on real text: 72% evidence recall on the untuned 18F handbook"], size=13, gap=4)
rows = [["Gate (Theme 4)", "Target", "Test-1", "Clean test-2"],
        ["G1 one-command run, unattended replay", "pass", "pass", "pass"],
        ["G2 early retrieval", "≥ 80%", "100%", "100%"],
        ["G3 multi-intent identification", "≥ 70%", "100%", "100%"],
        ["G4 citation support / fabricated", "≥ 85%", "100% / 0", "100% / 0"],
        ["G5 refinement continuity", "100%", "100%", "100%"],
        ["G6 telemetry coverage", "100%", "100%", "100%"],
        ["Evidence recall (baseline 62.7%)", "–", "89.5%", "95.8%"],
        ["Key-fact recall (baseline 62.3%)", "–", "85.5%", "83.3%"]]
table(s, Inches(6.75), Inches(1.5), Inches(5.65), rows, [Inches(2.75), Inches(0.85), Inches(1.0), Inches(1.05)], size=12, row_h=Inches(0.47))
text(s, Inches(6.75), Inches(1.5) + Inches(0.47) * 9 + Inches(0.1), Inches(5.65), Inches(0.5),
     "Test-2 was frozen before any error analysis and never used for tuning. Full report: reports/metrics.md.", size=11, color=MUTED)

# ---------------------------------------------------------------- slide 9: what's next
s = S[8]
drop_body(s)
nxt = [("On-device for Galaxy", "Port the controller and index to Android (ONNX / Kotlin); the 16 MB embedding model and ~1 ms decisions already fit a phone budget."),
       ("Audio-native endpointing", "Use VAD and prosody, not only transcript punctuation, to time retrieval and handle barge-in during the spoken answer."),
       ("Close the synonym gap", "A small cross-encoder, or the optional LLM, used only as a verifier when coverage is borderline (≈20–50 ms)."),
       ("Learn the controller", "Train the wait/retrieve policy on logged traces to trade lead time against wasted queries per user population."),
       ("Multilingual", "Hindi, Korean and code-mixed speech with multilingual static embeddings and language-aware segmentation."),
       ("Salient entities on real corpora", "Filter link text and boilerplate before entities can scope a query (the main 18F failure mode).")]
cw_ = (CW - Inches(0.3)) / 2
for i, (h_, b_) in enumerate(nxt):
    col, row = i % 2, i // 2
    x = LM + col * (cw_ + Inches(0.3))
    y = Inches(1.5) + row * Inches(1.62)
    card(s, x, y, cw_, Inches(1.45))
    text(s, x + Inches(0.25), y + Inches(0.15), cw_ - Inches(0.5), Inches(1.2),
         [[(h_, {"bold": True, "size": 16, "color": PURPLE})], [(b_, {"size": 13})]])

# ---------------------------------------------------------------- slide 10: differentiation
s = S[9]
drop_body(s)
diff = [("Reproduces the guide, line by line", "All three worked examples in the Theme 4 guide replay as specified: provisional retrieval at 0.8 s, refinement without restart, reformatting with retrieval_required: false."),
        ("Honest measurement", "Frozen splits with SHA-256 hashes, a clean test-2, a real untuned corpus, the first pre-analysis run kept, and every failure listed."),
        ("Nothing hard-coded", "Entities, scopes, aspects and capacity words are mined from the corpus; the same code runs on the 18F handbook."),
        ("Latency hiding, quantified", "With a 150 ms retrieval stack, Kairos starts answering in 2.4 ms against 152 ms for batch RAG (replayed traces)."),
        ("Zero-cost by default, LLM-ready", "CPU only, $0 per turn; an optional LLM is wrapped in a verifier that drops fabricated citations (tested)."),
        ("Production hygiene", "One-command Docker, CI building the image and running the replay suite, schema-validated traces, 16 tests.")]
for i, (h_, b_) in enumerate(diff):
    col, row = i % 3, i // 3
    cw3 = (CW - Inches(0.6)) / 3
    x = LM + col * (cw3 + Inches(0.3))
    y = Inches(1.5) + row * Inches(2.45)
    card(s, x, y, cw3, Inches(2.25))
    text(s, x + Inches(0.22), y + Inches(0.18), cw3 - Inches(0.44), Inches(2.0),
         [[(h_, {"bold": True, "size": 16, "color": PURPLE})], [(b_, {"size": 13})]])

# ---------------------------------------------------------------- slide 11: checklist
s = S[10]
body = next(sh for sh in s.placeholders if sh.placeholder_format.idx == 1)
checks = ["Working prototype code — public GitHub repo: Y (github.com/CodeThread386/kairos-streaming-rag)",
          "README with reproducible setup instructions: Y (one-command Docker run, local steps, API and CLI examples)",
          "Demo video, max 5 minutes: Y (demo/kairos_demo.mp4, 3 min 30 s, in the tagged commit)",
          "Presentation file (PPT or PDF): Y (presentation/VITV_Kairos_Submission.pptx + .pdf)",
          "AI disclosure: Y (docs/AI_DISCLOSURE.md + filled LangAI3.0 form)",
          "requirements.txt, Dockerfile, docker-compose.yml, tests, benchmark report: Y",
          "Release tag PRISM_GENAI_HACKATHON_Y2026 on the final commit: Y"]
txBody = body.text_frame._txBody
paras = body.text_frame.paragraphs
proto = copy.deepcopy(paras[0]._p)
for p in list(paras):
    txBody.remove(p._p)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
for c in checks:
    np_ = copy.deepcopy(proto)
    rs = np_.findall(A + "r")
    for extra in rs[1:]:
        np_.remove(extra)
    rs[0].find(A + "t").text = c
    txBody.append(np_)

prs.save(OUT)
print("saved", OUT)
